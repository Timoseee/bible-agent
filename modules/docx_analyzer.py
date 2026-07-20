"""DOCX template inspection utilities."""

from pathlib import Path

from docx import Document


def _length(value):
    return None if value is None else value.pt


def _color(value):
    try:
        rgb = value.rgb
    except AttributeError:
        return None

    return None if rgb is None else str(rgb)


def _alignment(value):
    return None if value is None else str(value)


def _line_spacing(value):
    if value is None:
        return None

    if hasattr(value, "pt"):
        return value.pt

    return value


def analyze_template(template_path):
    """Extract structure information from a DOCX template without modifying it."""
    path = Path(template_path)

    if not path.exists():
        raise FileNotFoundError(f"Template file not found: {path}")

    document = Document(path)
    styles = []
    fonts = set()
    colors = set()
    paragraph_formats = []

    for style in document.styles:
        styles.append(style.name)
        font = getattr(style, "font", None)

        if font is not None:
            if font.name:
                fonts.add(font.name)

            color = _color(font.color)
            if color:
                colors.add(color)

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else None
        paragraph_format = paragraph.paragraph_format
        paragraph_formats.append(
            {
                "style": style_name,
                "alignment": _alignment(paragraph.alignment),
                "line_spacing": _line_spacing(paragraph_format.line_spacing),
                "left_indent": _length(paragraph_format.left_indent),
                "right_indent": _length(paragraph_format.right_indent),
                "first_line_indent": _length(paragraph_format.first_line_indent),
                "space_before": _length(paragraph_format.space_before),
                "space_after": _length(paragraph_format.space_after),
            }
        )

        for run in paragraph.runs:
            if run.font.name:
                fonts.add(run.font.name)

            color = _color(run.font.color)
            if color:
                colors.add(color)

    sections = []
    for section in document.sections:
        sections.append(
            {
                "header_paragraphs": len(section.header.paragraphs),
                "footer_paragraphs": len(section.footer.paragraphs),
                "page_width": _length(section.page_width),
                "page_height": _length(section.page_height),
                "top_margin": _length(section.top_margin),
                "bottom_margin": _length(section.bottom_margin),
                "left_margin": _length(section.left_margin),
                "right_margin": _length(section.right_margin),
            }
        )

    return {
        "filename": path.name,
        "paragraph_count": len(document.paragraphs),
        "style_count": len(document.styles),
        "styles": styles,
        "fonts": sorted(fonts),
        "colors": sorted(colors),
        "paragraph_formats": paragraph_formats,
        "tables": len(document.tables),
        "sections": sections,
        "headers": sum(section["header_paragraphs"] for section in sections),
        "footers": sum(section["footer_paragraphs"] for section in sections),
    }
