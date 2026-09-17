"""Audio sermon DOCX generation using audio template formatting."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import RGBColor

from modules.paragraph_optimizer import structure_audio_paragraphs
from modules.text_cleaning_pipeline import to_simplified_chinese


def _capture_paragraph_format(paragraph):
    """Capture paragraph and first-run formatting from a template paragraph."""
    run = paragraph.runs[0] if paragraph.runs else None
    paragraph_format = paragraph.paragraph_format
    return {
        "style": paragraph.style,
        "alignment": paragraph.alignment,
        "first_line_indent": paragraph_format.first_line_indent,
        "line_spacing": paragraph_format.line_spacing,
        "space_before": paragraph_format.space_before,
        "space_after": paragraph_format.space_after,
        "font_name": run.font.name if run else None,
        "font_size": run.font.size if run else None,
        "bold": run.bold if run else None,
    }


def _extract_template_formats(document):
    """Use the first template paragraph as title format and the next as body format."""
    paragraphs = list(document.paragraphs)
    if not paragraphs:
        return None, None

    title_format = _capture_paragraph_format(paragraphs[0])
    body_source = paragraphs[1] if len(paragraphs) > 1 else paragraphs[0]
    body_format = _capture_paragraph_format(body_source)
    return title_format, body_format


def _clear_document_body(document):
    """Remove template sample paragraphs while keeping section settings."""
    body = document.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def _apply_run_format(run, format_info):
    if format_info.get("font_name"):
        run.font.name = format_info["font_name"]
        run._element.rPr.rFonts.set(qn("w:eastAsia"), format_info["font_name"])
    if format_info.get("font_size") is not None:
        run.font.size = format_info["font_size"]
    if format_info.get("bold") is not None:
        run.bold = format_info["bold"]
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.highlight_color = WD_COLOR_INDEX.AUTO


def _apply_paragraph_format(paragraph, format_info):
    if not format_info:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.highlight_color = WD_COLOR_INDEX.AUTO
        return

    paragraph_format = paragraph.paragraph_format
    paragraph.alignment = format_info.get("alignment")
    paragraph_format.first_line_indent = format_info.get("first_line_indent")
    paragraph_format.line_spacing = format_info.get("line_spacing")
    paragraph_format.space_before = format_info.get("space_before")
    paragraph_format.space_after = format_info.get("space_after")

    for run in paragraph.runs:
        _apply_run_format(run, format_info)


def _add_formatted_paragraph(document, text, format_info):
    style = format_info.get("style") if format_info else None
    try:
        paragraph = document.add_paragraph(style=style)
    except (KeyError, ValueError, AttributeError):
        paragraph = document.add_paragraph()

    paragraph.add_run(text)
    _apply_paragraph_format(paragraph, format_info)
    return paragraph


def generate_audio_docx(text, template_path, output_path):
    """Generate an audio sermon DOCX from corrected text and the audio template."""
    template = Path(template_path)
    output = Path(output_path)

    if not template.exists():
        raise FileNotFoundError(f"Template file not found: {template}")

    structured = structure_audio_paragraphs(to_simplified_chinese(text))
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document(template)
    title_format, body_format = _extract_template_formats(document)
    _clear_document_body(document)

    for item in structured:
        format_info = title_format if item["role"] == "title" else body_format
        _add_formatted_paragraph(document, item["text"], format_info)

    document.save(output)

    if not output.exists():
        raise RuntimeError(f"DOCX output was not created: {output}")

    Document(output)

    return {
        "output_path": str(output),
        "success": True,
        "paragraphs_inserted": len(structured),
    }
