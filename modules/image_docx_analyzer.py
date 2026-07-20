"""Deep analysis utilities for image DOCX templates."""

from pathlib import Path

from docx import Document


def _length(value):
    return None if value is None else value.pt


def _color(font):
    if font is None:
        return None

    try:
        rgb = font.color.rgb
    except AttributeError:
        return None

    return None if rgb is None else str(rgb)


def _alignment(value):
    return None if value is None else str(value)


def _line_spacing(value):
    if value is None:
        return None

    if hasattr(value, "pt"):
        return value.pt

    return value


def _paragraph_format(paragraph):
    paragraph_format = paragraph.paragraph_format
    return {
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": _alignment(paragraph.alignment),
        "line_spacing": _line_spacing(paragraph_format.line_spacing),
        "left_indent": _length(paragraph_format.left_indent),
        "right_indent": _length(paragraph_format.right_indent),
        "first_line_indent": _length(paragraph_format.first_line_indent),
        "space_before": _length(paragraph_format.space_before),
        "space_after": _length(paragraph_format.space_after),
    }


def _run_format(run):
    return {
        "text": run.text,
        "font": run.font.name,
        "size": _length(run.font.size),
        "bold": run.bold,
        "italic": run.italic,
        "underline": run.underline,
        "color": _color(run.font),
    }


def analyze_image_template(template_path):
    """Analyze the image DOCX template without modifying it."""
    path = Path(template_path)

    if not path.exists():
        raise FileNotFoundError(f"Template file not found: {path}")

    document = Document(path)
    styles = []
    runs = []
    colors = set()
    fonts = set()
    format_patterns = []

    for style in document.styles:
        font = getattr(style, "font", None)
        style_info = {
            "name": style.name,
            "type": str(style.type),
            "font": font.name if font is not None else None,
            "size": _length(font.size) if font is not None else None,
            "bold": font.bold if font is not None else None,
            "italic": font.italic if font is not None else None,
            "color": _color(font),
        }
        styles.append(style_info)

        if style_info["font"]:
            fonts.add(style_info["font"])
        if style_info["color"]:
            colors.add(style_info["color"])

    for paragraph_index, paragraph in enumerate(document.paragraphs):
        paragraph_info = _paragraph_format(paragraph)
        paragraph_info["paragraph_index"] = paragraph_index
        format_patterns.append(paragraph_info)

        for run_index, run in enumerate(paragraph.runs):
            run_info = _run_format(run)
            run_info["paragraph_index"] = paragraph_index
            run_info["run_index"] = run_index
            runs.append(run_info)

            if run_info["font"]:
                fonts.add(run_info["font"])
            if run_info["color"]:
                colors.add(run_info["color"])

    sections = []
    for section in document.sections:
        sections.append(
            {
                "page_width": _length(section.page_width),
                "page_height": _length(section.page_height),
                "top_margin": _length(section.top_margin),
                "bottom_margin": _length(section.bottom_margin),
                "left_margin": _length(section.left_margin),
                "right_margin": _length(section.right_margin),
                "header_paragraphs": len(section.header.paragraphs),
                "footer_paragraphs": len(section.footer.paragraphs),
            }
        )

    return {
        "filename": path.name,
        "paragraph_count": len(document.paragraphs),
        "styles": styles,
        "runs": runs,
        "colors": sorted(colors),
        "fonts": sorted(fonts),
        "format_patterns": format_patterns,
        "tables": len(document.tables),
        "sections": sections,
    }
