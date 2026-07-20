"""DOCX generation foundation that preserves template structure."""

from pathlib import Path

from docx import Document

from modules.audio_docx_formatter import generate_audio_docx


def generate_docx(text, template_path, output_path):
    """Load a template, insert plain text, save a new DOCX, and verify it."""
    template = Path(template_path)
    output = Path(output_path)

    if not template.exists():
        raise FileNotFoundError(f"Template file not found: {template}")

    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document(template)

    for paragraph_text in text.split("\n\n"):
        if paragraph_text.strip():
            document.add_paragraph(paragraph_text.strip())

    document.save(output)

    if not output.exists():
        raise RuntimeError(f"DOCX output was not created: {output}")

    Document(output)

    return {
        "output_path": str(output),
        "success": True,
        "paragraphs_inserted": len([part for part in text.split("\n\n") if part.strip()]),
    }


def generate_audio_document(text, template_path, output_path):
    """Generate an audio sermon DOCX using audio-specific formatting."""
    return generate_audio_docx(text, template_path, output_path)
