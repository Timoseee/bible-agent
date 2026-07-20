"""Phase 7B tests for audio DOCX generation."""

import hashlib
import unittest
from pathlib import Path

from docx import Document

from modules.audio_docx_formatter import generate_audio_docx
from modules.paragraph_optimizer import optimize_audio_paragraphs
from modules.template_manager import load_template


BASE_DIR = Path(__file__).resolve().parent.parent
AUDIO_TEMPLATE = BASE_DIR / "templates" / "audio_template.docx"
TEST_OUTPUT = BASE_DIR / "output" / "docx" / "phase7b_audio_test.docx"


def _file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Phase7BAudioDocxTests(unittest.TestCase):
    def tearDown(self):
        if TEST_OUTPUT.exists():
            TEST_OUTPUT.unlink()

    def test_audio_template_loads(self):
        template = load_template("audio")
        self.assertFalse(isinstance(template, str))

    def test_text_generates_valid_docx(self):
        result = generate_audio_docx("第一句。第二句。", AUDIO_TEMPLATE, TEST_OUTPUT)
        self.assertTrue(result["success"])
        self.assertTrue(TEST_OUTPUT.exists())

    def test_generated_docx_can_reopen(self):
        generate_audio_docx("第一句。第二句。", AUDIO_TEMPLATE, TEST_OUTPUT)
        document = Document(TEST_OUTPUT)
        self.assertGreater(len(document.paragraphs), 0)

    def test_empty_paragraphs_removed(self):
        result = generate_audio_docx("第一句。\n\n\n\n第二句。", AUDIO_TEMPLATE, TEST_OUTPUT)
        self.assertEqual(result["paragraphs_inserted"], 1)

    def test_font_color_becomes_black(self):
        generate_audio_docx("黑色文字测试。", AUDIO_TEMPLATE, TEST_OUTPUT)
        document = Document(TEST_OUTPUT)
        inserted_runs = [
            run
            for paragraph in document.paragraphs
            for run in paragraph.runs
            if "黑色文字测试" in run.text
        ]
        self.assertTrue(inserted_runs)
        self.assertEqual(str(inserted_runs[0].font.color.rgb), "000000")

    def test_original_template_unchanged(self):
        before = _file_hash(AUDIO_TEMPLATE)
        generate_audio_docx("模板不应被修改。", AUDIO_TEMPLATE, TEST_OUTPUT)
        after = _file_hash(AUDIO_TEMPLATE)
        self.assertEqual(before, after)

    def test_paragraph_optimization_works(self):
        text = "第一句。\n第二句。\n第三句。\n第四句。"
        paragraphs = optimize_audio_paragraphs(text)
        self.assertEqual(paragraphs, ["第一句。第二句。第三句。第四句。"])


if __name__ == "__main__":
    unittest.main()
