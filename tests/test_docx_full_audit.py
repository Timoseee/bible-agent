"""Tests for auditing an existing black paragraph DOCX."""

import json
import shutil
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Inches, RGBColor

from modules.docx_full_audit import audit_existing_docx


TEST_DIR = Path(__file__).resolve().parent / "tmp_docx_audit"


class DocxAuditProvider:
    def __init__(self):
        self.audit_calls = 0

    def generate_json(self, prompt, content):
        if "FULL_TEXT_AUDIT_REVIEW" in prompt:
            records = json.loads(content)["candidates"]
            return json.dumps(
                {
                    "decisions": [
                        {
                            "id": item["id"],
                            "approved": True,
                            "verified_replacement": item["replacement"],
                        }
                        for item in records
                    ]
                }
            )
        self.audit_calls += 1
        if self.audit_calls == 1:
            return json.dumps(
                {
                    "candidates": [
                        {
                            "paragraph_id": "P0001",
                            "original": "模西",
                            "replacement": "摩西",
                            "error_type": "bible_name",
                            "reason": "圣经人名错字",
                            "confidence": 0.99,
                        }
                    ]
                },
                ensure_ascii=False,
            )
        return '{"candidates": []}'


class ExistingDocxAuditTests(unittest.TestCase):
    def setUp(self):
        TEST_DIR.mkdir(parents=True, exist_ok=True)

    def tearDown(self):
        shutil.rmtree(TEST_DIR, ignore_errors=True)

    def test_preserves_paragraph_count_margins_and_black_text(self):
        source = TEST_DIR / "source.docx"
        output = TEST_DIR / "output.docx"
        document = Document()
        section = document.sections[0]
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        run = document.add_paragraph("模西上山。").runs[0]
        run.font.color.rgb = RGBColor(0, 0, 0)
        document.add_paragraph("第二段。")
        document.save(source)

        result = audit_existing_docx(source, output, DocxAuditProvider())
        verified = Document(output)
        self.assertEqual([p.text for p in verified.paragraphs], ["摩西上山。", "第二段。"])
        self.assertEqual(result["full_audit_applied"], 1)
        self.assertEqual(verified.sections[0].left_margin, verified.sections[0].right_margin)
        self.assertTrue(
            all(
                run.font.color.rgb == RGBColor(0, 0, 0)
                for paragraph in verified.paragraphs
                for run in paragraph.runs
                if run.text
            )
        )


if __name__ == "__main__":
    unittest.main()
