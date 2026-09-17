"""Phase 7D tests for final pipeline routing and DOCX rendering."""

import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from docx import Document

from modules.image_docx_renderer import generate_image_docx
from modules.pipeline_controller import process_input, process_audio_input
from modules.style_mapper import build_style_mapping
from modules.template_manager import get_template_path


BASE_DIR = Path(__file__).resolve().parent.parent
TEST_DIR = Path(tempfile.gettempdir()) / "BibleCorrectionAI_tmp_phase7d"
OUTPUT_DIR = BASE_DIR / "output" / "docx"


class MockProvider:
    def generate(self, prompt, content):
        if "Return JSON only" in prompt:
            return '{"approved": true, "issues": [], "summary": "Approved."}'
        if content.startswith("Original text:\n"):
            return content.split("\n", 1)[1].split("\n\n", 1)[0]
        if "Sermon text:\n" in content:
            return content.split("Sermon text:\n", 1)[1]
        return content

    def generate_json(self, prompt, content):
        _ = content
        if "FULL_TEXT_AUDIT_REVIEW" in prompt:
            return '{"decisions": []}'
        if "FULL_TEXT_AUDIT" in prompt:
            return '{"candidates": []}'
        return '{"approved": true, "issues": [], "summary": "Approved."}'

    def generate_from_image(self, prompt, image_path):
        _ = (prompt, image_path)
        return "图像OCR文本。"

    def analyze_image(self, image_path):
        return {
            "filename": Path(image_path).name,
            "page_number": 1,
            "chapter_reference": "",
            "first_text": "开始",
            "last_text": "结束",
            "confidence": 0.95,
        }


def _create_image(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (20, 20), color="white").save(path)


class Phase7DFinalPipelineTests(unittest.TestCase):
    def setUp(self):
        TEST_DIR.mkdir(parents=True, exist_ok=True)
        self.output_patcher = patch(
            "modules.pipeline_controller.OUTPUT_DOCX_DIR", TEST_DIR / "pipeline_output"
        )
        self.output_patcher.start()

    def tearDown(self):
        self.output_patcher.stop()
        if TEST_DIR.exists():
            shutil.rmtree(TEST_DIR)
        for output in OUTPUT_DIR.glob("audio_*.docx"):
            if "mock" in output.stem:
                output.unlink()
        for output in OUTPUT_DIR.glob("image_*.docx"):
            if "images" in output.stem or "page" in output.stem:
                output.unlink()
        test_docx = OUTPUT_DIR / "phase7d_image_render_test.docx"
        if test_docx.exists():
            test_docx.unlink()

    def test_audio_input_detection(self):
        audio_path = TEST_DIR / "mock_audio.mp3"
        audio_path.touch()
        with patch("modules.pipeline_controller.process_audio_input") as mock_audio:
            mock_audio.return_value = {"success": True, "input_type": "audio"}
            result = process_input(audio_path, provider=MockProvider())
        self.assertTrue(result["success"])
        mock_audio.assert_called_once()

    def test_image_folder_detection(self):
        image_folder = TEST_DIR / "images"
        _create_image(image_folder / "page.png")
        with patch("modules.pipeline_controller.process_image_input") as mock_image:
            mock_image.return_value = {"success": True, "input_type": "image"}
            result = process_input(
                image_folder,
                provider=MockProvider(),
                vision_provider=MockProvider(),
                cleanup_images=False,
            )
        self.assertTrue(result["success"])
        mock_image.assert_called_once()

    def test_pipeline_routing_for_unsupported_input(self):
        text_path = TEST_DIR / "notes.txt"
        text_path.write_text("unsupported", encoding="utf-8")
        result = process_input(text_path, provider=MockProvider())
        self.assertFalse(result["success"])
        self.assertEqual(result["step"], "input_detection")

    def test_mock_audio_workflow(self):
        audio_path = TEST_DIR / "mock_audio.mp3"
        audio_path.touch()
        with patch("modules.pipeline_controller.transcribe_audio") as mock_transcribe:
            mock_transcribe.return_value = {"filename": "mock_audio.mp3", "language": "zh", "text": "讲章文本。"}
            result = process_input(audio_path, provider=MockProvider())
        self.assertTrue(result["success"])
        self.assertEqual(result["input_type"], "audio")
        self.assertTrue(Path(result["output"]).exists())

    def test_mock_image_workflow(self):
        image_folder = TEST_DIR / "images"
        _create_image(image_folder / "page.png")
        result = process_input(
            image_folder,
            provider=MockProvider(),
            vision_provider=MockProvider(),
            cleanup_images=False,
        )
        self.assertTrue(result["success"])
        self.assertEqual(result["input_type"], "image")
        self.assertTrue(Path(result["output"]).exists())

    def test_audio_correction_failure_or_rejection_does_not_export(self):
        for review in ("Correction failed: connection failed", "Review rejected changes"):
            with self.subTest(review=review), patch(
                "modules.pipeline_controller.transcribe_audio", return_value={"text": "原始转录。"}
            ), patch(
                "modules.pipeline_controller.correct_with_bible_check",
                return_value={"approved": False, "final_text": "原始转录。", "review": review},
            ), patch("modules.pipeline_controller.generate_audio_docx") as render, patch(
                "modules.pipeline_controller.polish_sermon_text"
            ) as polish:
                result = process_audio_input("mock.m4a", provider=MockProvider())
                self.assertFalse(result["success"])
                self.assertEqual(result["step"], "correction")
                self.assertIn(review, result["error"])
                render.assert_not_called()
                polish.assert_not_called()

    def test_audio_polish_failure_does_not_export(self):
        with patch(
            "modules.pipeline_controller.transcribe_audio", return_value={"text": "讲章文本。"}
        ), patch(
            "modules.pipeline_controller.polish_sermon_text",
            return_value={"success": False, "polished_text": "讲章文本。", "warnings": ["connection failed"]},
        ), patch("modules.pipeline_controller.generate_audio_docx") as render:
            result = process_audio_input("mock.m4a", provider=MockProvider())
        self.assertFalse(result["success"])
        self.assertEqual(result["step"], "polish")
        self.assertIn("connection failed", result["error"])
        render.assert_not_called()

    def test_traditional_audio_becomes_simplified_in_docx_and_filename(self):
        with patch(
            "modules.pipeline_controller.transcribe_audio",
            return_value={"text": "民數記第四章，利未人看守約櫃。"},
        ), patch(
            "modules.pipeline_controller.polish_sermon_text",
            return_value={"success": True, "polished_text": "民數記第四章\n\n利未人看守約櫃。"},
        ):
            result = process_audio_input("mock.m4a", provider=MockProvider())
        self.assertTrue(result["success"])
        text = "\n".join(p.text for p in Document(result["output"]).paragraphs)
        self.assertIn("民数记", text)
        self.assertIn("利未人看守约柜。", text)
        self.assertNotIn("數", text)
        self.assertIn("民数记", Path(result["output"]).name)

    def test_docx_output_validation(self):
        output_path = OUTPUT_DIR / "phase7d_image_render_test.docx"
        mapping = build_style_mapping(get_template_path("image"))
        source = "标题\n\n出埃及记24章\n正文内容第一行，\n正文内容第二行。"
        result = generate_image_docx(source, mapping, get_template_path("image"), output_path)
        self.assertTrue(result["success"])
        document = Document(output_path)
        self.assertGreater(len(document.paragraphs), 0)

        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertEqual(document_text, "标题\n出埃及记24章正文内容第一行，正文内容第二行。")
        self.assertNotEqual(document.paragraphs[0].text, "出埃及记24章1/18节《组员回应补充》")
        section = document.sections[0]
        self.assertAlmostEqual(section.left_margin.inches, section.right_margin.inches)
        self.assertTrue(
            all(
                run.font.color.rgb is not None and str(run.font.color.rgb) == "000000"
                for paragraph in document.paragraphs
                for run in paragraph.runs
                if run.text
            )
        )

    def test_image_workflow_output_contains_reviewed_text(self):
        image_folder = TEST_DIR / "images"
        _create_image(image_folder / "page.png")
        result = process_input(
            image_folder,
            provider=MockProvider(),
            vision_provider=MockProvider(),
            cleanup_images=False,
        )

        self.assertTrue(result["success"])
        document = Document(result["output"])
        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("图像OCR文本。", document_text)
        self.assertNotIn("[IMAGE_001]", document_text)

    def test_image_ocr_text_is_rendered_without_placeholders(self):
        image_folder = TEST_DIR / "images"
        _create_image(image_folder / "page.png")
        ocr_text = "出埃及记第二十四章\n摩西上山"
        with patch("modules.pipeline_controller.extract_text_from_images") as mock_ocr:
            mock_ocr.return_value = {
                "text": ocr_text,
                "images_processed": 1,
                "results": [{"success": True, "text": ocr_text}],
            }
            result = process_input(
                image_folder,
                provider=MockProvider(),
                vision_provider=MockProvider(),
                cleanup_images=False,
            )

        self.assertTrue(result["success"])
        document = Document(result["output"])
        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("出埃及记第二十四章", document_text)
        self.assertIn("摩西上山", document_text)
        self.assertNotIn("[IMAGE_001]", document_text)

    def test_error_handling(self):
        audio_path = TEST_DIR / "mock_audio.mp3"
        audio_path.touch()
        with patch("modules.pipeline_controller.transcribe_audio", side_effect=RuntimeError("boom")):
            result = process_input(audio_path, provider=MockProvider())
        self.assertFalse(result["success"])
        self.assertEqual(result["step"], "transcription")
        self.assertIn("boom", result["error"])


if __name__ == "__main__":
    unittest.main()
